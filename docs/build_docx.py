#!/usr/bin/env python
"""把 docs/实验记录.md 转成 实验记录.docx（无 pandoc 依赖）。"""

from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "docs" / "实验记录.md"
TARGET = ROOT / "实验记录.docx"

INLINE_RE = re.compile(r"(\*\*.+?\*\*|`[^`]+`|\[[^\]]+\]\([^)]+\))")
HEADING_RE = re.compile(r"^(#{1,4})\s+(.*)$")
BULLET_RE = re.compile(r"^\s*[-*]\s+(.*)$")
NUMBER_RE = re.compile(r"^\s*\d+\.\s+(.*)$")
LINK_RE = re.compile(r"\[([^\]]+)\]\(([^)]+)\)")
IMAGE_RE = re.compile(r"^!\[([^\]]*)\]\(([^)]+)\)$")


def set_run_font(run, ascii_font: str = "Calibri", east: str = "宋体", size=None) -> None:
    run.font.name = ascii_font
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    rfonts.set(qn("w:eastAsia"), east)
    if size:
        run.font.size = Pt(size)


def add_inline(paragraph, text: str) -> None:
    for part in INLINE_RE.split(text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            run = paragraph.add_run(part[2:-2])
            run.bold = True
            set_run_font(run)
        elif part.startswith("`") and part.endswith("`"):
            run = paragraph.add_run(part[1:-1])
            set_run_font(run, "Consolas", "宋体", 9)
            run.font.color.rgb = RGBColor(0x33, 0x33, 0x33)
        elif part.startswith("["):
            match = LINK_RE.match(part)
            if match:
                label, url = match.groups()
                run = paragraph.add_run(label)
                set_run_font(run)
                note = paragraph.add_run(f"（{url}）")
                set_run_font(note, size=8)
                note.font.color.rgb = RGBColor(0x66, 0x66, 0x66)
            else:
                paragraph.add_run(part)
        else:
            run = paragraph.add_run(part)
            set_run_font(run)


def shade(paragraph, fill: str = "F2F2F2") -> None:
    ppr = paragraph._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    ppr.append(shd)


def add_code_block(doc: Document, lines: list[str]) -> None:
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(2)
    paragraph.paragraph_format.space_after = Pt(6)
    paragraph.paragraph_format.left_indent = Cm(0.3)
    shade(paragraph)
    run = paragraph.add_run("\n".join(lines))
    set_run_font(run, "Consolas", "宋体", 9)


def add_table(doc: Document, raw_lines: list[str]) -> None:
    rows: list[list[str]] = []
    for line in raw_lines:
        cells = [cell.strip() for cell in line.strip().strip("|").split("|")]
        if all(re.fullmatch(r":?-{3,}:?", cell) for cell in cells if cell):
            continue
        rows.append(cells)
    if not rows:
        return
    columns = max(len(row) for row in rows)
    table = doc.add_table(rows=len(rows), cols=columns)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for r_index, row in enumerate(rows):
        for c_index in range(columns):
            cell = table.cell(r_index, c_index)
            cell.text = ""
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.space_after = Pt(0)
            text = row[c_index] if c_index < len(row) else ""
            add_inline(paragraph, text)
            for run in paragraph.runs:
                set_run_font(run, size=9)
                if r_index == 0:
                    run.bold = True
    doc.add_paragraph()


def add_image(doc: Document, caption: str, rel_path: str, number: int) -> int:
    """插入图片并加“图 N”题注；路径相对 docs/ 解析。"""
    path = (SOURCE.parent / rel_path).resolve()
    if not path.exists():
        paragraph = doc.add_paragraph()
        run = paragraph.add_run(f"[图片缺失] {rel_path}")
        set_run_font(run, size=9)
        run.font.color.rgb = RGBColor(0xC0, 0x00, 0x00)
        return number

    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(6)
    paragraph.paragraph_format.space_after = Pt(2)
    paragraph.add_run().add_picture(str(path), width=Cm(15.5))

    number += 1
    caption_paragraph = doc.add_paragraph()
    caption_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption_paragraph.paragraph_format.space_after = Pt(8)
    run = caption_paragraph.add_run(f"图 {number}  {caption}")
    set_run_font(run, size=9)
    run.font.color.rgb = RGBColor(0x55, 0x55, 0x55)
    return number


def main() -> None:
    lines = SOURCE.read_text(encoding="utf-8").splitlines()
    doc = Document()

    section = doc.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.2)
    section.bottom_margin = Cm(2.2)
    section.left_margin = Cm(2.2)
    section.right_margin = Cm(2.2)

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10.5)
    normal.element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), "宋体")
    normal.paragraph_format.line_spacing = 1.25
    normal.paragraph_format.space_after = Pt(4)

    image_no = 0
    index = 0
    while index < len(lines):
        line = lines[index]

        if line.startswith("```"):
            index += 1
            code: list[str] = []
            while index < len(lines) and not lines[index].startswith("```"):
                code.append(lines[index])
                index += 1
            index += 1
            add_code_block(doc, code)
            continue

        if line.strip().startswith("|"):
            table_lines: list[str] = []
            while index < len(lines) and lines[index].strip().startswith("|"):
                table_lines.append(lines[index])
                index += 1
            add_table(doc, table_lines)
            continue

        image = IMAGE_RE.match(line.strip())
        if image:
            image_no = add_image(doc, image.group(1), image.group(2), image_no)
            index += 1
            continue

        heading = HEADING_RE.match(line)
        if heading:
            level = len(heading.group(1))
            paragraph = doc.add_heading(level=min(level, 4))
            paragraph.paragraph_format.space_before = Pt(10 if level <= 2 else 6)
            paragraph.paragraph_format.space_after = Pt(4)
            add_inline(paragraph, heading.group(2))
            for run in paragraph.runs:
                set_run_font(run, "Calibri", "黑体", 16 if level == 1 else 13 if level == 2 else 11.5)
                run.font.color.rgb = RGBColor(0x1F, 0x1F, 0x1F)
                run.bold = True
            index += 1
            continue

        bullet = BULLET_RE.match(line)
        if bullet:
            paragraph = doc.add_paragraph(style="List Bullet")
            add_inline(paragraph, bullet.group(1))
            index += 1
            continue

        numbered = NUMBER_RE.match(line)
        if numbered:
            paragraph = doc.add_paragraph(style="List Number")
            add_inline(paragraph, numbered.group(1))
            index += 1
            continue

        if line.startswith(">"):
            quote_lines = []
            while index < len(lines) and lines[index].startswith(">"):
                quote_lines.append(lines[index].lstrip(">").strip())
                index += 1
            paragraph = doc.add_paragraph()
            paragraph.paragraph_format.left_indent = Cm(0.6)
            add_inline(paragraph, " ".join(quote_lines))
            for run in paragraph.runs:
                run.italic = True
                run.font.color.rgb = RGBColor(0x55, 0x55, 0x55)
            continue

        if not line.strip():
            index += 1
            continue

        # 合并连续普通行，避免 Markdown 软换行被拆成多段
        paragraph_lines = [line.strip()]
        index += 1
        while index < len(lines):
            nxt = lines[index]
            if (
                not nxt.strip()
                or nxt.startswith(("#", "```", ">", "|"))
                or BULLET_RE.match(nxt)
                or NUMBER_RE.match(nxt)
            ):
                break
            paragraph_lines.append(nxt.strip())
            index += 1
        paragraph = doc.add_paragraph()
        add_inline(paragraph, " ".join(paragraph_lines))

    doc.save(TARGET)
    print(f"[完成] 已生成 {TARGET}")


if __name__ == "__main__":
    main()
